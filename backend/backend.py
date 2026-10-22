import re
import sqlite3
from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt
from fpdf import FPDF

DB_PATH = Path(__file__).parent / "data" / "database.db"

LABELS = {
    "vyst_mo": {
        "codvuz": "Код вуза",
        "type": "Форма НИР",
        "regnumber": "Рег. №",
        "exhitype": "Экспонат",
        "subject": "Наименование НИР",
        "grnti": "Код ГРНТИ",
        "shortname": "Вуз",
        "bossname": "Руководитель",
        "bosstitle": "Должность",
        "vystavki": "Информация о выставке",
        "exponat": "Название экспоната",
    },
    "vuz": {
        "codvuz": "Код",
        "shortname": "Сокр. наименование",
        "city": "Город",
        "profile": "Профиль",
        "status": "Статус",
        "gr_ved": "Тип",
        "name": "Наименование",
        "fullname": "Полное наименование",
        "region": "Федеральный округ",
        "obl": "Код субъекта",
        "oblname": "Субъект РФ",
    },
    "grntirub": {"codrub": "Код", "rubrika": "Рубрика"},
}
KEY = ["codvuz", "type", "regnumber"]
NIR_TYPES = ["Е", "М"]
EXHIBIT_TYPES = {"Е": "есть", "П": "планируется", "Н": "нет"}


def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def query(sql, params=()):
    with connect() as conn:
        return pd.read_sql_query(sql, conn, params=params)


def get_table(table):
    return query(f"SELECT * FROM {table}")


def sort_by_key(df, ascending=True):
    tmp = df.assign(
        _vuz=pd.to_numeric(df["codvuz"], errors="coerce"),
        _reg=pd.to_numeric(df["regnumber"], errors="coerce"),
    )
    order = tmp.sort_values(["_vuz", "type", "_reg", "regnumber"], ascending=ascending)
    return df.loc[order.index].reset_index(drop=True)


def split_grnti(value):
    return [p for p in re.split(r"[;,\s]+", str(value or "")) if p]


GEO_FIELDS = ("region", "oblname", "city", "codvuz")


def vuz_names():
    vuz = get_table("vuz")
    return dict(zip(vuz["codvuz"], vuz["shortname"]))


def rubric_names():
    rub = get_table("grntirub")
    return {c: f"{c:02d} - {r}" for c, r in zip(rub["codrub"], rub["rubrika"])}


def record_rubrics(grnti):
    return {int(c[:2]) for c in split_grnti(grnti) if re.match(r"\d{2}\.", c)}


def apply_filter(df, flt):
    geo = {k: v for k, v in flt.items() if k in GEO_FIELDS and v}
    if geo:
        vuz = get_table("vuz")
        for field, values in geo.items():
            vuz = vuz[vuz[field].isin(values)]
        df = df[df["codvuz"].isin(vuz["codvuz"])]
    if flt.get("rubrics"):
        df = df[
            df["grnti"].apply(lambda v: bool(record_rubrics(v) & set(flt["rubrics"])))
        ]
    if flt.get("exhitype"):
        df = df[df["exhitype"].isin(flt["exhitype"])]
    return df


def describe_filter(flt):
    labels = {
        "region": "Федеральный округ",
        "oblname": "Субъект РФ",
        "city": "Город",
        "codvuz": "Вуз",
        "rubrics": "Рубрика ГРНТИ",
        "exhitype": "Экспонат",
    }
    parts = []
    for key, value in flt.items():
        if key == "exhitype":
            value = [EXHIBIT_TYPES[v] for v in value]
        if key == "codvuz":
            names = vuz_names()
            value = [names[v] for v in value]
        if key == "rubrics":
            names = rubric_names()
            value = [names[v] for v in value]
        parts.append(f"{labels[key]}: {', '.join(value)}")
    return "; ".join(parts) or "без фильтра"


def validate(rec, rec_id=None):
    errors = [
        ((f,), f"Не заполнено поле «{LABELS['vyst_mo'][f]}»")
        for f in ("codvuz", "type", "regnumber", "subject", "grnti", "exhitype")
        if not rec.get(f)
    ]
    required = {"Е": ("vystavki", "exponat"), "П": ("exponat",)}.get(
        rec.get("exhitype"), ()
    )
    errors += [
        (
            ("exhitype", f),
            f"Экспонат {EXHIBIT_TYPES[rec['exhitype']]}: заполните поле «{LABELS['vyst_mo'][f]}»",
        )
        for f in required
        if not rec.get(f)
    ]
    if errors:
        return errors
    with connect() as conn:
        if not conn.execute(
            "SELECT 1 FROM vuz WHERE codvuz = ?", (rec["codvuz"],)
        ).fetchone():
            errors.append((("codvuz",), "Вуза с таким кодом нет в справочнике"))
        if conn.execute(
            "SELECT 1 FROM vyst_mo WHERE codvuz = ? AND type = ? AND regnumber = ? AND id IS NOT ?",
            (rec["codvuz"], rec["type"], rec["regnumber"], rec_id),
        ).fetchone():
            errors.append(
                (tuple(KEY), "Запись с таким ключом (вуз + форма + рег. №) уже есть")
            )
        rubrics = {r for (r,) in conn.execute("SELECT codrub FROM grntirub")}
    for code in split_grnti(rec["grnti"]):
        if not re.fullmatch(r"\d{2}\.\d{2}(\.\d{2})?", code):
            errors.append((("grnti",), f"Код ГРНТИ «{code}» должен иметь вид ХХ.ХХ.ХХ"))
        elif int(code[:2]) not in rubrics:
            errors.append((("grnti",), f"Рубрики ГРНТИ {code[:2]} нет в справочнике"))
    return errors


def save_record(rec, rec_id=None):
    rec = {
        k: None if v is None or v != v else str(v).strip() or None for k, v in rec.items()
    }
    old = {}
    if rec_id is not None:
        old = query("SELECT * FROM vyst_mo WHERE id = ?", (rec_id,)).iloc[0].to_dict()
        old = {k: None if pd.isna(v) else str(v).strip() or None for k, v in old.items()}
    if rec.get("grnti"):
        rec["grnti"] = "; ".join(split_grnti(rec["grnti"]))
    errors = [
        (fields, msg)
        for fields, msg in validate(rec, rec_id)
        if rec_id is None or any(rec.get(f) != old.get(f) for f in fields)
    ]
    if errors:
        return errors, None
    with connect() as conn:
        row = conn.execute(
            "SELECT shortname FROM vuz WHERE codvuz = ?", (rec["codvuz"],)
        ).fetchone()
        rec["shortname"] = row[0] if row else old.get("shortname")
        if rec_id is None:
            rec_id = conn.execute(
                f"INSERT INTO vyst_mo ({', '.join(rec)}) VALUES ({', '.join('?' * len(rec))})",
                list(rec.values()),
            ).lastrowid
        else:
            conn.execute(
                f"UPDATE vyst_mo SET {', '.join(f'{k} = ?' for k in rec)} WHERE id = ?",
                [*rec.values(), rec_id],
            )
    return [], rec_id


def execute(sql, params=()):
    with connect() as conn:
        conn.execute(sql, params)


def nir_card(rec_id):
    r = query(
        "SELECT v.*, z.fullname, z.city, z.oblname, z.region FROM vyst_mo v "
        "LEFT JOIN vuz z ON z.codvuz = v.codvuz WHERE v.id = ?",
        (rec_id,),
    ).iloc[0]
    r = {k: None if pd.isna(v) else v for k, v in r.items()}
    rubrics = rubric_names()
    grnti = "; ".join(
        (
            f"{c} — {rubrics[int(c[:2])].split(' - ', 1)[1]}"
            if re.match(r"\d{2}\.", c) and int(c[:2]) in rubrics
            else c
        )
        for c in split_grnti(r["grnti"])
    )
    nir_type = {"Е": "тематический план", "М": "научно-техническая программа"}
    return [
        ("Вуз", f"{r['codvuz']} — {r['shortname'] or ''}"),
        ("Полное наименование вуза", r["fullname"]),
        ("Город", r["city"]),
        ("Субъект РФ", r["oblname"]),
        ("Федеральный округ", r["region"]),
        ("Форма НИР", f"{r['type']} — {nir_type.get(r['type'], '')}"),
        ("Регистрационный номер", r["regnumber"]),
        ("Наименование НИР", r["subject"]),
        ("Коды ГРНТИ", grnti),
        ("Руководитель", r["bossname"]),
        ("Должность руководителя", r["bosstitle"]),
        ("Выставочный экспонат", EXHIBIT_TYPES.get(r["exhitype"], r["exhitype"])),
        ("Информация о выставке", r["vystavki"]),
        ("Название экспоната", r["exponat"]),
    ]


FONTS = Path(__file__).parent / "fonts"
PDF_WIDTHS = {
    "Значение": 4,
    "Наименование НИР": 4,
    "Информация о выставке": 4,
    "Название экспоната": 4,
    "Руководитель": 2.4,
    "Должность": 2,
    "Код ГРНТИ": 1.6,
    "Вуз": 1.6,
    "Экспонат": 1.8,
    "Форма НИР": 1.2,
}
FORMATS = {
    "PDF": ("pdf", "application/pdf"),
    "Word": (
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    "Excel": (
        "xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    "CSV": ("csv", "text/csv"),
}


def to_pdf(df, title, notes):
    wide = len(df.columns) > 3
    pdf = FPDF(orientation="L" if wide else "P")
    pdf.add_font("DejaVu", fname=FONTS / "DejaVuSans.ttf")
    pdf.add_font("DejaVu", style="B", fname=FONTS / "DejaVuSans-Bold.ttf")
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 14)
    pdf.multi_cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", size=9)
    for note in notes:
        pdf.multi_cell(0, 6, note, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("DejaVu", size=7 if wide else 10)
    widths = [PDF_WIDTHS.get(c, 1) for c in df]
    with pdf.table(
        col_widths=widths, text_align="LEFT", line_height=4 if wide else 6
    ) as table:
        for row in [list(df.columns), *df.astype(str).values.tolist()]:
            cells = table.row()
            for value in row:
                cells.cell(value or "—")
    return bytes(pdf.output())


def to_docx(df, title, notes):
    doc = Document()
    section = doc.sections[0]
    if len(df.columns) > 3:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    doc.add_heading(title, level=1)
    for note in notes:
        doc.add_paragraph(note)
    table = doc.add_table(rows=1, cols=len(df.columns), style="Table Grid")
    size = Pt(7 if len(df.columns) > 3 else 10)
    for cell, name in zip(table.rows[0].cells, df.columns):
        run = cell.paragraphs[0].add_run(str(name))
        run.bold, run.font.size = True, size
    for row in df.astype(str).values.tolist():
        for cell, value in zip(table.add_row().cells, row):
            cell.paragraphs[0].add_run(value or "—").font.size = size
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export(df, fmt, title, notes=()):
    ext, mime = FORMATS[fmt]
    df = df.fillna("")
    if ext == "csv":
        data = df.to_csv(index=False).encode("utf-8-sig")
    elif ext == "xlsx":
        buf = BytesIO()
        df.to_excel(buf, index=False, sheet_name=re.sub(r"[:\\/?*\[\]]", " ", title)[:31])
        data = buf.getvalue()
    elif ext == "pdf":
        data = to_pdf(df, title, notes)
    else:
        data = to_docx(df, title, notes)
    return data, ext, mime


def get_groups():
    return query(
        "SELECT g.id, g.name, g.filter_desc, COUNT(i.nir_id) AS n "
        "FROM nir_groups g LEFT JOIN nir_group_items i ON i.group_id = g.id "
        "GROUP BY g.id ORDER BY g.name"
    )


def add_to_group(name, filter_desc, nir_ids):
    with connect() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO nir_groups (name, filter_desc) VALUES (?, ?)",
            (name, filter_desc),
        )
        (group_id,) = conn.execute(
            "SELECT id FROM nir_groups WHERE name = ?", (name,)
        ).fetchone()
        conn.executemany(
            "INSERT OR IGNORE INTO nir_group_items (group_id, nir_id) VALUES (?, ?)",
            [(group_id, int(i)) for i in nir_ids],
        )


def get_group_items(group_id):
    df = query(
        "SELECT v.* FROM nir_group_items i JOIN vyst_mo v ON v.id = i.nir_id "
        "WHERE i.group_id = ?",
        (group_id,),
    )
    return sort_by_key(df)


def merge_groups(source_ids, name, filter_desc):
    with connect() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO nir_groups (name, filter_desc) VALUES (?, ?)",
            (name, filter_desc),
        )
        (target,) = conn.execute(
            "SELECT id FROM nir_groups WHERE name = ?", (name,)
        ).fetchone()
        conn.execute(
            "INSERT OR IGNORE INTO nir_group_items (group_id, nir_id) "
            "SELECT DISTINCT ?, nir_id FROM nir_group_items "
            f"WHERE group_id IN ({', '.join('?' * len(source_ids))}) AND group_id != ?",
            (target, *map(int, source_ids), target),
        )
    return target


def remove_from_group(group_id, nir_ids):
    with connect() as conn:
        conn.executemany(
            "DELETE FROM nir_group_items WHERE group_id = ? AND nir_id = ?",
            [(group_id, int(i)) for i in nir_ids],
        )
